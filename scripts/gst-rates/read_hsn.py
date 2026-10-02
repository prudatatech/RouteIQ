"""3. Read the GST portal HSN list (HSN_SAC.xlsx, sheet HSN_MSTR) from $GST_WORK into hsn_rows.json (needs openpyxl)."""
import sys
import openpyxl, os, json, collections
WORK = os.environ.get('GST_WORK') or sys.exit('Set GST_WORK to the folder that holds the downloads (never commit them)')
os.chdir(WORK)
wb = openpyxl.load_workbook('HSN_SAC.xlsx', read_only=True)
print(wb.sheetnames)
ws = wb['HSN_MSTR']
rows = []
for i, r in enumerate(ws.iter_rows(values_only=True)):
    if i == 0:
        print('header', r)
        continue
    if r[0] is None:
        continue
    rows.append([str(r[0]).strip(), (r[1] or '').strip() if isinstance(r[1], str) else str(r[1])])
print(len(rows), collections.Counter(len(c) for c, _ in rows))
print(rows[:5], rows[-3:])
bad = [c for c, _ in rows if not c.isdigit()]
print('non-digit', bad[:20], len(bad))
dups = [c for c, n in collections.Counter(c for c, _ in rows).items() if n > 1]
print('dups', len(dups), dups[:10])
json.dump(rows, open('hsn_rows.json', 'w'))
for c, d in rows:
    if c in ('1006', '10063010', '2523', '3401', '8517', '2202', '22029990', '22029920', '22029930', '22029991', '6904', '0910'):
        print(c, d[:120])
